"""Editable Word templates: A4 letterhead and proposal/report document."""
import os
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(__file__))
from brand import EMAILS, HEX, PLACEHOLDER as PH, WEBSITE  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'out')
LOGO = os.path.join(OUT, '01-Logos')
IMG = os.path.join(ROOT, 'build', 'office')


def C(name):
    return RGBColor.from_string(HEX[name].lstrip('#'))


RED, OBS, ASH, DEEP, STONE = C('Signal Red'), C('Obsidian'), C('Ash'), C('Deep Red'), C('Stone')
DISPLAY, BODY, MONO = 'Space Grotesk', 'Inter', 'JetBrains Mono'
MARKETS = 'SG · VN · ID · MY · TH'
HAIR = 'DDDCD7'


def set_font(run_or_style, name, size=None, color=None, bold=None):
    f = run_or_style.font
    f.name = name
    rpr = run_or_style.element.get_or_add_rPr() if hasattr(run_or_style, 'element') and run_or_style.element.tag.endswith('}r') else run_or_style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(a), name)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if fonts.get(qn(a)) is not None:
            del fonts.attrib[qn(a)]
    if size:
        f.size = Pt(size)
    if color is not None:
        f.color.rgb = color
    if bold is not None:
        f.bold = bold


def run(p, text, name=BODY, size=10, color=OBS, bold=False, spacing=None):
    r = p.add_run(text)
    set_font(r, name, size, color, bold)
    if spacing:
        sp = OxmlElement('w:spacing')
        sp.set(qn('w:val'), str(int(spacing)))
        r.element.get_or_add_rPr().append(sp)
    return r


def border(p, side='top', color=HAIR, size=4, space=6):
    ppr = p._p.get_or_add_pPr()
    bdr = ppr.find(qn('w:pBdr'))
    if bdr is None:
        bdr = OxmlElement('w:pBdr')
        ppr.append(bdr)
    el = OxmlElement(f'w:{side}')
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), str(size))
    el.set(qn('w:space'), str(space))
    el.set(qn('w:color'), color)
    bdr.append(el)


def shade(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcpr.append(shd)


def no_borders(table):
    tblpr = table._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{side}')
        e.set(qn('w:val'), 'nil')
        b.append(e)
    tblpr.append(b)


def fixed_widths(table, widths):
    table.autofit = False
    grid = table._tbl.tblGrid
    for col, w in zip(grid.findall(qn('w:gridCol')), widths):
        col.set(qn('w:w'), str(int(w / 635)))  # EMU to twentieths of a point
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w


def split_row(container, width=Mm(170)):
    t = container.add_table(rows=1, cols=2, width=width)
    no_borders(t)
    fixed_widths(t, [Mm(120), width - Mm(120)])
    left, right = t.rows[0].cells
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return left.paragraphs[0], right.paragraphs[0]


def cell_border(cell, side, color, size):
    tcpr = cell._tc.get_or_add_tcPr()
    b = tcpr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders')
        tcpr.append(b)
    e = OxmlElement(f'w:{side}')
    e.set(qn('w:val'), 'single')
    e.set(qn('w:sz'), str(size))
    e.set(qn('w:color'), color)
    b.append(e)


def page_field(p, size=7.5, color=ASH):
    """Insert a PAGE number field."""
    r = run(p, '', MONO, size, color)
    for kind, text in (('begin', None), (None, 'PAGE'), ('end', None)):
        if kind:
            fc = OxmlElement('w:fldChar')
            fc.set(qn('w:fldCharType'), kind)
            r._r.append(fc)
        else:
            it = OxmlElement('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = text
            r._r.append(it)


def base_styles(doc):
    st = doc.styles
    normal = st['Normal']
    set_font(normal, BODY, 10, OBS)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.3
    for name, size, before, after in (('Heading 1', 20, 18, 6), ('Heading 2', 13.5, 14, 4), ('Title', 30, 0, 6)):
        s = st[name]
        set_font(s, DISPLAY, size, OBS, True)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
        # remove theme colours and borders Word adds to Title
        ppr = s.element.get_or_add_pPr()
        for b in ppr.findall(qn('w:pBdr')):
            ppr.remove(b)
    lb = st['List Bullet']
    set_font(lb, BODY, 10, OBS)


def a4(section, top, bottom, left=20, right=20, header=10, footer=10):
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(top), Mm(bottom)
    section.left_margin, section.right_margin = Mm(left), Mm(right)
    section.header_distance, section.footer_distance = Mm(header), Mm(footer)


def contact_footer(section):
    f = section.footer
    rule = f.paragraphs[0]
    border(rule, 'bottom', HAIR, 4, 1)
    rule.paragraph_format.space_after = Pt(4)
    left, right = split_row(f)
    run(left, f"{PH['address'][0]} · {PH['address'][1]}, {PH['address'][2]}", BODY, 6.8, ASH)
    left.paragraph_format.space_after = Pt(0)
    l2 = left._parent.add_paragraph()
    run(l2, PH['reg'], BODY, 6.8, ASH)
    run(right, '■ ', BODY, 6.5, RED)
    run(right, MARKETS, MONO, 6.5, ASH, spacing=20)


# ------------------------------------------------------------------ letterhead
def letterhead():
    doc = Document()
    base_styles(doc)
    sec = doc.sections[0]
    a4(sec, 48, 30, header=16, footer=12)
    h = sec.header
    t = h.add_table(rows=1, cols=2, width=Mm(170))
    no_borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.columns[0].width, t.columns[1].width = Mm(100), Mm(70)
    left, right = t.rows[0].cells
    left.width, right.width = Mm(100), Mm(70)
    lp = left.paragraphs[0]
    lp.add_run().add_picture(os.path.join(LOGO, 'Horizontal', 'png', 'logic-sonata-horizontal-on-light-3000px.png'), width=Mm(58))
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for i, line in enumerate([WEBSITE, EMAILS['sales'], PH['phone']]):
        if i:
            rp = right.add_paragraph()
            rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rp.paragraph_format.space_after = Pt(0)
        run(rp, line, MONO, 7, ASH)
    h.paragraphs[0].text = ''
    rule = h.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.add_run().add_picture(os.path.join(IMG, 'red-rule.png'), width=Mm(12))
    contact_footer(sec)

    body = [
        ('30 September 2026', {}),
        ('', {}),
        ('Recipient Name\nJob Title\nCompany Name\nStreet Address\nCity Postcode', {}),
        ('', {}),
        ('Subject line in one short sentence', {'bold': True, 'font': DISPLAY, 'size': 12}),
        ('Dear [Recipient Name],', {}),
        ('Thank you for your time on [date]. This letter template uses Inter at 10 pt for body copy with generous line spacing. Keep letters to one page where possible and write in plain, direct sentences.', {}),
        ('Replace this text with your message. The header and footer carry the logo, contact details and company registration, so they never need to be retyped. Choose File > Save As > Word Template (.dotx) to keep a clean master.', {}),
        ('Yours sincerely,', {}),
        ('', {}),
        (PH['name'], {'bold': True, 'font': DISPLAY, 'size': 11}),
        (f"{PH['title']}, Logic Sonata", {'color': ASH}),
    ]
    for text, o in body:
        p = doc.add_paragraph()
        run(p, text, o.get('font', BODY), o.get('size', 10), o.get('color', OBS), o.get('bold', False))
    path = os.path.join(OUT, '04-Stationery', 'Letterhead', 'logic-sonata-letterhead.docx')
    doc.save(path)
    return path


# ------------------------------------------------------------------ proposal / report
def proposal():
    doc = Document()
    base_styles(doc)
    cover = doc.sections[0]
    a4(cover, 0, 12, left=0, right=0, header=0, footer=8)
    cover.different_first_page_header_footer = True
    p = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(os.path.join(IMG, 'proposal-cover-band.png'), width=Mm(210))

    def cover_par(text, font, size, color, bold=False, before=0, after=4, spacing=None):
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Mm(20)
        q.paragraph_format.right_indent = Mm(20)
        q.paragraph_format.space_before = Pt(before)
        q.paragraph_format.space_after = Pt(after)
        q.paragraph_format.line_spacing = 1.05
        run(q, text, font, size, color, bold, spacing)
        return q

    cover_par('■  PROPOSAL', MONO, 9, RED, before=40, after=10, spacing=40)
    cover_par('Proposal title in one or two lines', DISPLAY, 32, OBS, True, after=8)
    cover_par('A one-sentence subtitle that states the outcome for the client.', BODY, 13, ASH, after=36)
    for label, value in [('PREPARED FOR', 'Client Company Name'), ('PREPARED BY', f"{PH['name']}, {PH['title']}"), ('DATE', '30 September 2026'), ('REFERENCE', 'LS-2026-000')]:
        cover_par(label, MONO, 7.5, ASH, after=1, spacing=30)
        cover_par(value, BODY, 11, OBS, after=10)
    conf = cover_par('Confidential. Prepared for the named recipient only.', BODY, 8, ASH, before=60)
    run(conf, f'    {WEBSITE}', MONO, 8, DEEP)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    a4(body, 30, 24, header=12, footer=10)
    body.different_first_page_header_footer = False
    body.header.is_linked_to_previous = False
    body.footer.is_linked_to_previous = False
    hl, hr = split_row(body.header)
    hl.add_run().add_picture(os.path.join(LOGO, 'Horizontal', 'png', 'logic-sonata-horizontal-on-light-1000px.png'), width=Mm(38))
    run(hr, 'PROPOSAL TITLE', MONO, 7, ASH, spacing=20)
    rule = body.footer.paragraphs[0]
    border(rule, 'bottom', HAIR, 4, 1)
    rule.paragraph_format.space_after = Pt(4)
    fl, fr = split_row(body.footer)
    run(fl, 'Logic Sonata · Confidential', BODY, 7.5, ASH)
    page_field(fr)

    doc.add_heading('1. Executive summary', 1)
    doc.add_paragraph('Open with the client’s situation and the outcome this proposal delivers, in three or four sentences. Body text is Inter 10 pt; headings are Space Grotesk. Keep paragraphs short and specific.')
    doc.add_heading('What we heard', 2)
    for b in ['The client’s first priority, in their words', 'The constraint that shapes the solution (data, budget, timing)', 'The measure of success both sides agree on']:
        doc.add_paragraph(b, style='List Bullet')
    doc.add_heading('2. Proposed approach', 1)
    doc.add_paragraph('Describe the solution and how it will be deployed (on-premise, hosted or hybrid). Refer to products by name and code, for example Private Knowledge Assistant (PAI-KB).')

    call = doc.add_table(rows=1, cols=1)
    no_borders(call)
    c = call.rows[0].cells[0]
    shade(c, 'F2F1EE')
    cp = c.paragraphs[0]
    cell_border(c, 'left', 'FF3B30', 24)
    fixed_widths(call, [Mm(170)])
    run(cp, 'KEY POINT  ', MONO, 7.5, DEEP, spacing=30)
    run(cp, 'Use a callout for the one idea the reader must not miss.', BODY, 10, OBS)
    doc.add_paragraph()

    doc.add_heading('3. Scope and timeline', 1)
    rows = [('Phase', 'Deliverable', 'Timing'), ('Assess', 'Readiness report and 30/60/90-day roadmap', 'Weeks 1 to 2'),
            ('Pilot', 'One working use case, proven with real users', 'Weeks 3 to 8'), ('Roll out', 'Department deployment, training and governance', 'Weeks 9 to 16'),
            ('Manage', 'Managed support tier and quarterly reviews', 'Ongoing')]
    t = doc.add_table(rows=len(rows), cols=3)
    no_borders(t)
    widths = [Mm(32), Mm(100), Mm(38)]
    fixed_widths(t, widths)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.width = widths[j]
            if i == 0:
                shade(cell, '0B0C11')
            elif i % 2 == 0:
                shade(cell, 'F2F1EE')
            q = cell.paragraphs[0]
            q.paragraph_format.space_after = Pt(0)
            if i == 0:
                run(q, val.upper(), MONO, 7.5, RGBColor(0xFF, 0xFF, 0xFF), spacing=20)
            else:
                run(q, val, MONO if j == 2 else BODY, 9.5 if j < 2 else 8.5, OBS if j < 2 else ASH, bold=(j == 0))
    doc.add_paragraph()
    doc.add_heading('4. Next steps', 1)
    doc.add_paragraph('State the decision you need, by when, and who to contact. End with one clear call to action.')
    p = doc.add_paragraph()
    run(p, f"{PH['name']}  ·  {PH['email']}  ·  {PH['phone']}", MONO, 9, DEEP)

    path = os.path.join(OUT, '06-Documents', 'logic-sonata-proposal-template.docx')
    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path)
    return path


if __name__ == '__main__':
    for f in (letterhead(), proposal()):
        print('wrote', os.path.relpath(f, ROOT))
