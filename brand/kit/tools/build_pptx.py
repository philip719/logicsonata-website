"""Editable PowerPoint files: presentation template, business cards, name badge, social posts."""
import os
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Mm, Pt

sys.path.insert(0, os.path.dirname(__file__))
from brand import EMAILS, HEX, PLACEHOLDER as PH, WEBSITE  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'out')
LOGO = os.path.join(OUT, '01-Logos')
IMG = os.path.join(ROOT, 'build', 'office')


def C(name):
    return RGBColor.from_string(HEX[name].lstrip('#'))


RED, OBS, PAPER, WHITE, GRAPHITE, SLATE, STONE, ASH, DEEP, AMBER, STEEL = (
    C('Signal Red'), C('Obsidian'), C('Paper'), C('White'), C('Graphite'), C('Slate'), C('Stone'), C('Ash'), C('Deep Red'), C('Amber'), C('Steel'))
DISPLAY, BODY, MONO = 'Space Grotesk', 'Inter', 'JetBrains Mono'
MARKETS = 'SG · VN · ID · MY · TH'


def logo_png(lockup, variant, px=3000):
    name = lockup.lower()
    size = {'Mark': 2048}.get(lockup, px)
    return os.path.join(LOGO, lockup, 'png', f'logic-sonata-{name}-{variant}-{size}px.png')


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    # Drop the theme's shape style so no default shadow or effect is applied anywhere.
    style = s._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if style is not None:
        s._element.remove(style)
    return s


def text(slide, x, y, w, h, runs, size, font=BODY, color=PAPER, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, line=None):
    """runs: a string, or a list of paragraphs, each a string or list of (text, color) pairs."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line:
            p.line_spacing = line
        parts = [(para, color)] if isinstance(para, str) else para
        for t, col in parts:
            r = p.add_run()
            r.text = t
            f = r.font
            f.size = Pt(size)
            f.name = font
            f.bold = bold
            f.color.rgb = col
            if spacing is not None:
                rPr = r._r.get_or_add_rPr()
                rPr.set('spc', str(int(spacing * 100)))
    return tb


def eyebrow(slide, x, y, label, color=STONE, marker=RED, size=11):
    rect(slide, x, y + Pt(size) * 0.25, Pt(size * 0.72), Pt(size * 0.72), marker)
    text(slide, x + Pt(size * 1.5), y, Mm(200), Pt(size * 1.6), label.upper(), size, MONO, color, spacing=2.2)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


# ------------------------------------------------------------------ presentation template
def deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9, 13.333 x 7.5 in
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]
    M = Mm(16)  # outer margin
    count = {'n': 0}

    def slide(bg=OBS, image=None):
        s = prs.slides.add_slide(blank)
        count['n'] += 1
        if image:
            s.shapes.add_picture(image, 0, 0, W, H)
        else:
            rect(s, 0, 0, W, H, bg)
        return s

    def footer(s, dark=True):
        s.shapes.add_picture(logo_png('Mark', 'on-dark' if dark else 'on-light'), M, H - Mm(13), height=Mm(5))
        text(s, M + Mm(12), H - Mm(13), Mm(120), Mm(6), 'LOGIC SONATA · CONFIDENTIAL', 9, MONO, STONE if dark else ASH, spacing=1.8, anchor=MSO_ANCHOR.MIDDLE)
        text(s, W - M - Mm(20), H - Mm(13), Mm(20), Mm(6), f'{count["n"]:02d}', 9, MONO, STONE if dark else ASH, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

    def title(s, eb, t, dark=True, y=Mm(18)):
        eyebrow(s, M, y, eb, STONE if dark else ASH)
        text(s, M, y + Mm(9), W - 2 * M, Mm(24), t, 34, DISPLAY, PAPER if dark else OBS, bold=True)

    # 1 Title
    s = slide(image=os.path.join(IMG, 'slide-bg-dark.png'))
    s.shapes.add_picture(logo_png('Horizontal', 'on-dark'), M, Mm(16), height=Mm(10))
    eyebrow(s, M, Mm(62), 'Presentation · Month 2026')
    text(s, M, Mm(72), Mm(250), Mm(60), [[('Presentation title', PAPER)], [('goes on two lines.', RED)]], 54, DISPLAY, bold=True, line=0.95)
    text(s, M, Mm(128), Mm(220), Mm(14), 'A one-line subtitle that sets up the conversation.', 20, BODY, STONE)
    text(s, M, H - Mm(24), Mm(150), Mm(12), [[(PH['name'], PAPER)], [(PH['title'], STONE)]], 12, BODY)
    text(s, W - M - Mm(80), H - Mm(19), Mm(80), Mm(8), MARKETS, 10, MONO, STONE, align=PP_ALIGN.RIGHT, spacing=2)
    notes(s, 'Title slide. Replace the title, subtitle, date and presenter. Keep the red on the last phrase of the title only.')

    # 2 Agenda
    s = slide()
    title(s, 'Agenda', 'What we will cover today.')
    items = ['The problem with public AI', 'Our private AI approach', 'Proposed pilot and timeline', 'Next steps']
    for i, it in enumerate(items):
        y = Mm(62) + i * Mm(22)
        rect(s, M, y, W - 2 * M, Emu(9525), SLATE)
        text(s, M, y + Mm(5), Mm(20), Mm(12), f'{i + 1:02d}', 18, MONO, RED)
        text(s, M + Mm(26), y + Mm(4), Mm(220), Mm(12), it, 24, DISPLAY, PAPER)
    footer(s)
    notes(s, 'Agenda. Up to six items; duplicate a row group for more.')

    # 3 Section divider
    s = slide(bg=RED)
    text(s, M, Mm(34), Mm(200), Mm(70), '01', 150, DISPLAY, WHITE, bold=True, line=0.9)
    text(s, M, Mm(100), Mm(260), Mm(40), 'Section title in two or three words.', 44, DISPLAY, WHITE, bold=True, line=1.0)
    s.shapes.add_picture(logo_png('Horizontal', 'white'), M, H - Mm(20), height=Mm(7))
    notes(s, 'Section divider on Signal Red. White text only. Use for at most one divider per section.')

    # 4 Title and content
    s = slide()
    title(s, 'Context', 'A clear slide title that states the point.')
    text(s, M, Mm(58), Mm(150), Mm(40), 'Lead paragraph in Inter, 20 pt. One or two sentences that frame the slide and tell the audience what to take away.', 20, BODY, PAPER, line=1.25)
    bullets = ['Supporting point one, kept to a single line', 'Supporting point two, with a concrete fact', 'Supporting point three, with the outcome']
    for i, b in enumerate(bullets):
        y = Mm(104) + i * Mm(12)
        rect(s, M, y + Mm(2.3), Mm(2.2), Mm(2.2), RED)
        text(s, M + Mm(7), y, Mm(150), Mm(10), b, 16, BODY, STONE)
    panel = rect(s, Mm(200), Mm(58), W - Mm(200) - M, Mm(88), GRAPHITE)
    text(s, Mm(208), Mm(66), Mm(110), Mm(10), 'KEY TAKEAWAY', 11, MONO, RED, spacing=2)
    text(s, Mm(208), Mm(78), Mm(110), Mm(60), 'Use this panel for the single most important number, quote or decision on the slide.', 20, DISPLAY, PAPER, line=1.15)
    footer(s)
    notes(s, 'Title and content. Keep body text to 40 words; move detail to the speaker notes or an appendix.')

    # 5 Two columns
    s = slide()
    title(s, 'Comparison', 'Two options, side by side.')
    for i, (h, b) in enumerate([('PUBLIC AI', 'Prompts, files and IP are processed on infrastructure you do not control, with no audit trail.'), ('PRIVATE AI', 'Model, knowledge base and access controls run inside your environment, fully audited.')]):
        x = M + i * ((W - 2 * M) / 2 + Mm(4))
        w = (W - 2 * M) / 2 - Mm(4)
        rect(s, x, Mm(60), w, Emu(19050), RED if i else STONE)
        text(s, x, Mm(66), w, Mm(8), h, 12, MONO, RED if i else STONE, spacing=2)
        text(s, x, Mm(78), w - Mm(10), Mm(60), b, 22, DISPLAY, PAPER, line=1.2)
    footer(s)

    # 6 Statement
    s = slide(bg=GRAPHITE)
    text(s, M, Mm(30), Mm(40), Mm(40), '“', 120, DISPLAY, RED, bold=True)
    text(s, M, Mm(62), Mm(280), Mm(60), 'A single statement or customer quote, set large, that the audience should remember.', 40, DISPLAY, PAPER, bold=True, line=1.1)
    rect(s, M, Mm(140), Mm(18), Emu(38100), RED)
    text(s, M, Mm(146), Mm(200), Mm(14), [[('Name Surname', PAPER)], [('Title, Company', STONE)]], 14, BODY)
    footer(s)

    # 7 Key figures
    s = slide()
    title(s, 'Impact', 'Three numbers that tell the story.')
    figs = [('100%', 'of data stays under your control'), ('3', 'deployment models: on-premise, hosted, hybrid'), ('5', 'Southeast Asian markets served')]
    cw = (W - 2 * M - Mm(16)) / 3
    for i, (v, l) in enumerate(figs):
        x = M + i * (cw + Mm(8))
        rect(s, x, Mm(66), cw, Emu(28575), RED)
        text(s, x, Mm(74), cw, Mm(36), v, 72, DISPLAY, PAPER, bold=True)
        text(s, x, Mm(112), cw - Mm(10), Mm(20), l, 16, BODY, STONE, line=1.2)
    footer(s)
    notes(s, 'Key figures. Only use real, sourced numbers. Keep to three.')

    # 8 Chart (native, editable)
    s = slide()
    title(s, 'Data', 'Chart title that states the finding.')
    cd = CategoryChartData()
    cd.categories = ['Q1', 'Q2', 'Q3', 'Q4']
    cd.add_series('Series A', (12, 19, 27, 34))
    cd.add_series('Series B', (8, 11, 15, 22))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, M, Mm(56), W - 2 * M - Mm(90), Mm(92), cd)
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    ch.legend.font.color.rgb = STONE
    ch.legend.font.size = Pt(12)
    ch.legend.font.name = BODY
    for ser, col in zip(ch.series, [RED, STONE]):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = col
    ch.plots[0].gap_width = 80
    for ax in (ch.category_axis, ch.value_axis):
        ax.tick_labels.font.color.rgb = STONE
        ax.tick_labels.font.size = Pt(12)
        ax.tick_labels.font.name = MONO
        ax.format.line.color.rgb = SLATE
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.major_gridlines.format.line.color.rgb = SLATE
    text(s, W - M - Mm(80), Mm(60), Mm(80), Mm(10), 'WHAT IT MEANS', 11, MONO, RED, spacing=2)
    text(s, W - M - Mm(80), Mm(72), Mm(80), Mm(60), 'One or two sentences explaining the insight behind the chart.', 18, DISPLAY, PAPER, line=1.2)
    text(s, M, Mm(152), Mm(200), Mm(8), 'Source: add the data source here.', 10, MONO, STONE)
    footer(s)
    notes(s, 'Native chart: right-click > Edit Data to change the numbers. Signal Red for the main series, Stone for comparison, Amber only for a flagged value.')

    # 9 Image and text
    s = slide()
    ph = rect(s, M, Mm(18), Mm(150), H - Mm(40), GRAPHITE)
    text(s, M, Mm(18), Mm(150), H - Mm(40), ['Drop an image here', 'Insert a picture and crop it to this frame'], 12, MONO, STONE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for (x, y, dx, dy) in [(M, Mm(18), 1, 1), (M + Mm(150), Mm(18), -1, 1), (M, H - Mm(22), 1, -1), (M + Mm(150), H - Mm(22), -1, -1)]:
        rect(s, x - (Mm(6) if dx < 0 else 0), y - (Emu(19050) if dy < 0 else 0), Mm(6), Emu(19050), RED)
        rect(s, x - (Emu(19050) if dx < 0 else 0), y - (Mm(6) if dy < 0 else 0), Emu(19050), Mm(6), RED)
    eyebrow(s, Mm(182), Mm(40), 'Case study')
    text(s, Mm(182), Mm(50), Mm(150), Mm(40), 'Headline about the image or case.', 32, DISPLAY, PAPER, bold=True, line=1.05)
    text(s, Mm(182), Mm(92), Mm(140), Mm(50), 'Two or three sentences of supporting copy. Photography should be real, well lit and uncluttered; avoid stock clichés.', 16, BODY, STONE, line=1.3)
    footer(s)

    # 10 Process
    s = slide()
    title(s, 'Process', 'From first conversation to managed operations.')
    steps = [('Assess', 'Review AI usage, data exposure and use cases.'), ('Pilot', 'Prove one use case with real users.'), ('Roll out', 'Expand with governance and training.'), ('Manage', 'Keep it supported and improving.')]
    sw = (W - 2 * M) / 4
    rect(s, M + Mm(7), Mm(76), W - 2 * M - Mm(40), Emu(12700), SLATE)
    for i, (h, b) in enumerate(steps):
        x = M + i * sw
        node = rect(s, x, Mm(69), Mm(14), Mm(14), GRAPHITE, MSO_SHAPE.OVAL)
        node.line.color.rgb = RED
        node.line.width = Pt(1.5)
        text(s, x, Mm(69), Mm(14), Mm(14), f'{i + 1:02d}', 12, MONO, RED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x, Mm(92), sw - Mm(8), Mm(10), h, 22, DISPLAY, PAPER, bold=True)
        text(s, x, Mm(104), sw - Mm(12), Mm(30), b, 14, BODY, STONE, line=1.25)
    footer(s)

    # 11 Table (native)
    s = slide()
    title(s, 'Plan', 'Scope, deliverables and timing.')
    rows = [('Phase', 'Deliverable', 'Timing'), ('Assess', 'Readiness report and 30/60/90-day roadmap', 'Weeks 1 to 2'), ('Pilot', 'One working use case with real users', 'Weeks 3 to 8'),
            ('Roll out', 'Department deployment, training, governance', 'Weeks 9 to 16'), ('Manage', 'Managed support and quarterly reviews', 'Ongoing')]
    tbl = s.shapes.add_table(len(rows), 3, M, Mm(58), W - 2 * M, Mm(80)).table
    # 'No Style, No Grid': the brand fills below do the work, without default borders.
    style = tbl._tbl.tblPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId')
    if style is None:
        from pptx.oxml.ns import qn
        style = tbl._tbl.tblPr.makeelement(qn('a:tableStyleId'), {})
        tbl._tbl.tblPr.append(style)
    style.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
    widths = [Mm(50), W - 2 * M - Mm(110), Mm(60)]
    for j, w in enumerate(widths):
        tbl.columns[j].width = w
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = SLATE if i == 0 else (GRAPHITE if i % 2 else OBS)
            cell.margin_left = cell.margin_right = Mm(4)
            tf = cell.text_frame
            tf.text = val
            r = tf.paragraphs[0].runs[0]
            r.font.size = Pt(12 if i == 0 else 15)
            r.font.name = MONO if i == 0 or j == 2 else BODY
            r.font.color.rgb = RED if (i == 0) else (PAPER if j < 2 else STONE)
    footer(s)
    notes(s, 'Native table: add rows with Tab from the last cell. Header row in Slate with Signal Red labels.')

    # 12 Light content slide
    s = slide(image=os.path.join(IMG, 'slide-bg-light.png'))
    title(s, 'Light layout', 'Use the light layout for print-heavy decks.', dark=False)
    text(s, M, Mm(60), Mm(200), Mm(40), 'The same structure on Paper, with Obsidian text. Signal Red stays for large type, rules and markers; use Deep Red for small red text on light grounds.', 20, BODY, OBS, line=1.3)
    for i, b in enumerate(['Obsidian headlines and body copy', 'Ash for secondary text', 'Signal Red for accents only']):
        y = Mm(104) + i * Mm(12)
        rect(s, M, y + Mm(2.3), Mm(2.2), Mm(2.2), RED)
        text(s, M + Mm(7), y, Mm(150), Mm(10), b, 16, BODY, ASH)
    footer(s, dark=False)

    # 13 Closing
    s = slide(image=os.path.join(IMG, 'slide-bg-dark.png'))
    s.shapes.add_picture(logo_png('Stacked-Tagline', 'on-dark'), M, Mm(24), height=Mm(40))
    text(s, M, Mm(82), Mm(250), Mm(30), [[('Thank ', PAPER), ('you.', RED)]], 60, DISPLAY, bold=True)
    text(s, M, Mm(114), Mm(200), Mm(30), [[(PH['name'], PAPER)], [(PH['title'], STONE)], [(f"{PH['email']}  ·  {PH['phone']}", STONE)]], 14, BODY, line=1.3)
    text(s, W - M - Mm(110), Mm(126), Mm(110), Mm(10), WEBSITE.upper(), 14, MONO, RED, align=PP_ALIGN.RIGHT, spacing=2)
    text(s, W - M - Mm(110), Mm(136), Mm(110), Mm(10), MARKETS, 11, MONO, STONE, align=PP_ALIGN.RIGHT, spacing=2)
    notes(s, 'Closing slide with contact details.')

    path = os.path.join(OUT, '05-Presentation', 'logic-sonata-presentation-template.pptx')
    os.makedirs(os.path.dirname(path), exist_ok=True)
    prs.save(path)
    return path


# ------------------------------------------------------------------ business cards (editable)
def cards():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Mm(96), Mm(60)
    blank = prs.slide_layouts[6]
    B = Mm(3)
    for theme in ['dark', 'light']:
        dark = theme == 'dark'
        # front
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(os.path.join(IMG, f'card-bg-{theme}.png'), 0, 0, prs.slide_width, prs.slide_height)
        lw = Mm(54)
        s.shapes.add_picture(logo_png('Horizontal', 'on-dark' if dark else 'on-light'), (prs.slide_width - lw) // 2, B + Mm(17), width=lw)
        text(s, 0, B + Mm(30.8), prs.slide_width, Mm(4), [[('INTELLIGENCE', STONE if dark else ASH), ('.', RED), (' HARMONY', STONE if dark else ASH), ('.', RED), (' IMPACT', STONE if dark else ASH), ('.', RED)]], 5.4, MONO, align=PP_ALIGN.CENTER, spacing=1.3)
        notes(s, f'{theme.title()} edition, front. Artboard includes 3 mm bleed on every side; the card trims to 90 x 54 mm. Export: File > Export > PDF (best quality).')
        # back
        s = prs.slides.add_slide(blank)
        rect(s, 0, 0, prs.slide_width, prs.slide_height, OBS if dark else WHITE)
        rect(s, 0, 0, B + Mm(1.6), prs.slide_height, RED)
        ink, muted, label = (PAPER, STONE, RED) if dark else (OBS, ASH, DEEP)
        L = B + Mm(7)
        s.shapes.add_picture(logo_png('Mark', 'on-dark' if dark else 'on-light'), prs.slide_width - B - Mm(6) - Mm(9.5), B + Mm(6), width=Mm(9.5))
        rect(s, L, B + Mm(6.9), Mm(1.5), Mm(1.5), RED)
        text(s, L + Mm(3), B + Mm(6.2), Mm(40), Mm(3), 'PRIVATE AI', 5.2, MONO, muted, spacing=1)
        text(s, L, B + Mm(12.3), Mm(60), Mm(6), PH['name'], 12.5, DISPLAY, ink, bold=True)
        text(s, L, B + Mm(19.2), Mm(60), Mm(4), PH['title'], 7, BODY, muted)
        rect(s, L, B + Mm(28), prs.slide_width - L - B - Mm(6), Emu(9000), SLATE if dark else RGBColor(0xDD, 0xDC, 0xD7))
        text(s, L, B + Mm(31), Mm(70), Mm(12), [[('T  ', label), (PH['phone'], ink)], [('E  ', label), (PH['email'], ink)], [('W  ', label), (WEBSITE, ink)]], 6.3, MONO, line=1.45)
        text(s, prs.slide_width - B - Mm(6) - Mm(40), prs.slide_height - B - Mm(8.5), Mm(40), Mm(3), MARKETS, 5, MONO, muted, align=PP_ALIGN.RIGHT, spacing=0.8)
        notes(s, f'{theme.title()} edition, back. Replace name, title, phone and email. Keep text 5 mm inside the trim.')
    path = os.path.join(OUT, '04-Stationery', 'Business-Cards', 'logic-sonata-business-cards-editable.pptx')
    prs.save(path)
    return path


# ------------------------------------------------------------------ name badge (editable)
def badge():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Mm(106), Mm(76)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    B = Mm(3)
    rect(s, 0, 0, prs.slide_width, prs.slide_height, WHITE)
    rect(s, 0, 0, prs.slide_width, B + Mm(24), OBS)
    rect(s, 0, B + Mm(24), prs.slide_width, Mm(1.2), RED)
    s.shapes.add_picture(logo_png('Horizontal', 'on-dark'), B + Mm(7), B + Mm(8.5), width=Mm(46))
    text(s, B + Mm(7), B + Mm(32), Mm(90), Mm(10), PH['name'], 19, DISPLAY, OBS, bold=True)
    text(s, B + Mm(7), B + Mm(43), Mm(90), Mm(6), PH['title'], 9.5, BODY, ASH)
    text(s, B + Mm(7), prs.slide_height - B - Mm(9), Mm(45), Mm(4), 'SINGAPORE', 7, MONO, DEEP, spacing=1.2)
    text(s, prs.slide_width - B - Mm(7) - Mm(50), prs.slide_height - B - Mm(9), Mm(50), Mm(4), WEBSITE, 7, MONO, ASH, align=PP_ALIGN.RIGHT, spacing=1)
    notes(s, 'Name badge, 100 x 70 mm plus 3 mm bleed. Duplicate the slide for each person, then export to PDF.')
    path = os.path.join(OUT, '08-Events', 'logic-sonata-name-badge-editable.pptx')
    prs.save(path)
    return path


# ------------------------------------------------------------------ social posts (editable, 1200 x 1200 px)
def social():
    prs = Presentation()
    px = Emu(9525)  # one pixel at 96 dpi
    prs.slide_width = prs.slide_height = px * 1200
    blank = prs.slide_layouts[6]

    def P(n):
        return px * n

    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, P(1200), P(1200), OBS)
    s.shapes.add_picture(logo_png('Horizontal', 'on-dark'), P(96), P(96), width=P(330))
    rect(s, P(96), P(428), P(14), P(14), RED)
    text(s, P(126), P(420), P(600), P(34), 'ANNOUNCEMENT', 18, MONO, STONE, spacing=4)
    text(s, P(96), P(480), P(1008), P(330), [[('Your headline goes ', PAPER), ('here.', RED)]], 78, DISPLAY, bold=True, line=0.98)
    text(s, P(96), P(830), P(900), P(150), 'One or two lines of supporting copy that explain the news and why it matters.', 25, BODY, STONE, line=1.35)
    text(s, P(96), P(1080), P(700), P(30), WEBSITE.upper(), 16, MONO, RED, spacing=3.5)
    notes(s, 'Announcement post, 1200 x 1200 px. Export: File > Export > PNG, current slide.')

    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, P(1200), P(1200), GRAPHITE)
    text(s, P(96), P(120), P(300), P(260), '“', 190, DISPLAY, RED, bold=True)
    text(s, P(96), P(380), P(1008), P(430), 'A customer or partner quote, in their own words, set large and kept short.', 49, DISPLAY, PAPER, bold=True, line=1.15)
    rect(s, P(96), P(860), P(80), P(6), RED)
    text(s, P(96), P(895), P(800), P(90), [[('Name Surname', PAPER)], [('Title, Company', STONE)]], 25, BODY, line=1.25)
    s.shapes.add_picture(logo_png('Mark', 'on-dark'), P(994), P(1031), width=P(110))
    notes(s, 'Quote post, 1200 x 1200 px. Always get written approval to quote a customer.')

    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, P(1200), P(1200), RED)
    rect(s, P(96), P(118), P(14), P(14), OBS)
    text(s, P(126), P(110), P(600), P(34), 'BY THE NUMBERS', 18, MONO, WHITE, spacing=4)
    text(s, P(88), P(250), P(1100), P(460), '100%', 300, DISPLAY, WHITE, bold=True)
    text(s, P(96), P(720), P(1008), P(200), 'of your data stays under your control.', 46, DISPLAY, OBS, bold=True, line=1.1)
    s.shapes.add_picture(logo_png('Horizontal', 'white'), P(96), P(1063), width=P(300))
    notes(s, 'Statistic post on Signal Red, 1200 x 1200 px. Use real, sourced numbers only.')

    path = os.path.join(OUT, '07-Social', 'logic-sonata-social-posts-editable.pptx')
    prs.save(path)
    return path


if __name__ == '__main__':
    for f in (deck(), cards(), badge(), social()):
        print('wrote', os.path.relpath(f, ROOT))
